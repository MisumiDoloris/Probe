"""端到端验证。对着 snmpsim 起的假设备跑一遍完整链路。

采集 -> 快照 -> 台账 -> Word 报告，第二次再跑一遍 -> diff。
"""

import sys

from probe.collect.base import Device
from probe.collect.snmp import SnmpCollector
from probe.report import snapshot, inventory, diff as diff_mod, doc_report

TARGET = "127.0.0.1"
PORT = 1161
COMMUNITY = "public"


def collect_once():
    dev = Device(ip=TARGET)
    try:
        data = SnmpCollector(COMMUNITY, port=PORT).collect(
            TARGET, {"community": COMMUNITY, "port": PORT}
        )
        for k, v in data.items():
            setattr(dev, k, v)
        dev.reachable = True
    except Exception as e:
        dev.error = f"{type(e).__name__}: {e}"
    print(f"采集 {TARGET}:{PORT}  {'成功' if dev.reachable else '失败'}")
    print(f"  设备名 {dev.hostname}")
    print(f"  版本   {dev.version[:70]}")
    return dev.to_dict()


def check_report(snap, diff_result=None, out="output/巡检报告_sim.docx"):
    """报告必须跟快照对得上。数字不一致就当失败，别拿眼睛看一遍说没事。"""
    from docx import Document

    doc_report.save(snap, out, diff_result=diff_result, meta={"org": "（单位名称）"})
    doc = Document(out)
    devices = snap.get("devices", [])

    stats_tbl = doc.tables[1]
    total = int(stats_tbl.rows[1].cells[0].text)
    failed = int(stats_tbl.rows[1].cells[2].text)
    real_failed = len([x for x in devices if x.get("error")])

    assert total == len(devices), f"报告设备数 {total}，快照 {len(devices)}，对不上"
    assert failed == real_failed, f"报告采集失败数 {failed}，快照 {real_failed}，对不上"

    ledger = doc.tables[2]
    rows = len(ledger.rows) - 1
    assert rows == len(devices), f"台账 {rows} 行，设备 {len(devices)} 台，对不上"

    print(f"报告自检通过：{total} 台，采集失败 {failed} 台，台账 {rows} 行 -> {out}")
    return out


def main():
    d = collect_once()
    path = snapshot.save([d], note="e2e")
    print(f"快照 {path}")
    excel = inventory.export_excel([d], "output/台账_sim.xlsx")
    print(f"台账 {excel}")

    snap = snapshot.load(path)

    if len(sys.argv) > 1 and sys.argv[1] == "diff":
        import os
        files = sorted(f for f in os.listdir("snapshots") if f.startswith("snapshot_"))
        if len(files) >= 2:
            old = snapshot.load(os.path.join("snapshots", files[-2]))
            new = snapshot.load(os.path.join("snapshots", files[-1]))
            result = diff_mod.compare(old, new)
            print("\n=== diff ===")
            for line in diff_mod.to_lines(result):
                print(line)
            print()
            check_report(snap, result)
            return

    check_report(snap)


if __name__ == "__main__":
    main()
